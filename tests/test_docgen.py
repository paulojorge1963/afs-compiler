"""Tests for the .docx generator (app.docgen), independent of the FastAPI layer.

Complements test_calc.py: those tests check the numbers are right, these check
that a valid report actually turns into a well-formed .docx, that a failing
validation blocks generation, and that notes/statement lines that don't apply
are cleanly omitted from the generated document (Section 9's acceptance tests).
"""
from docx import Document

from app.docgen import generate_afs_docx
from app.reporting import build_report


def _table_cell_texts(doc):
    return [cell.text for table in doc.tables for row in table.rows for cell in row.cells]


def test_generate_docx_for_full_example(example_financial_year):
    report = build_report(example_financial_year)
    assert report.is_valid

    buf = generate_afs_docx(report)
    doc = Document(buf)

    assert len(doc.tables) > 0
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "EXAMPLE PHOTOGRAPHY STUDIO (PTY) LTD" in all_text
    assert "FOR REVIEW BEFORE ISSUE" in all_text

    cells = _table_cell_texts(doc)
    assert any("23 447" in c for c in cells)  # profit for the year ties out


def test_generate_docx_blocked_when_invalid(example_financial_year):
    # Break the balance sheet by inflating a single trial balance line.
    example_financial_year.trial_balance_lines[0].current_year_amount += 1000
    report = build_report(example_financial_year)
    assert not report.is_valid

    try:
        generate_afs_docx(report)
        assert False, "expected generate_afs_docx to refuse an invalid report"
    except ValueError:
        pass


def test_generate_docx_omits_inapplicable_notes_for_minimal_entity(minimal_financial_year):
    report = build_report(minimal_financial_year)
    assert report.is_valid

    buf = generate_afs_docx(report)
    doc = Document(buf)
    all_text = "\n".join(p.text for p in doc.paragraphs)

    for heading in ["Property, plant and equipment", "Loans to (from) shareholders", "Inventories",
                     "Trade and other receivables", "Trade and other payables"]:
        assert heading not in all_text, f"unexpected note heading present: {heading}"
