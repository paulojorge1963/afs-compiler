"""Exports a FinancialYear back into the "AFS Compiler Input Template" workbook layout
(Section 4, user flow 4) - used by "Create Next Year" to hand back a workbook pre-filled
with opening balances and the prior-year comparative column already populated, and
generally as a take-the-data-offline convenience.

The sheet/column layout mirrors app.importer exactly, so a round trip re-imports cleanly.
"""
from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Font

from app.models import POLICY_AREAS
from app.reporting import ReportData

YELLOW = "FFFFF2CC"
BOLD = Font(bold=True)


def _label_value_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.cell(row=1, column=2, value=title).font = Font(bold=True, size=14)
    r = 3
    for row in rows:
        if row is None:
            r += 1
            continue
        if len(row) == 1:
            ws.cell(row=r, column=2, value=row[0]).font = BOLD
        else:
            label, value = row
            ws.cell(row=r, column=2, value=label)
            ws.cell(row=r, column=3, value=value)
        r += 1
    return ws


def export_workbook_for_year(fy, report: ReportData) -> BytesIO:
    entity = fy.entity
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _label_value_sheet(
        wb,
        "Entity Setup",
        [
            ("One workbook = one entity, one financial year. Yellow = fill in.",),
            None,
            ("COMPANY DETAILS",),
            ("Company Name", entity.company_name),
            ("Registration Number", entity.registration_number),
            ("Tax Reference Number", entity.tax_reference_number),
            ("VAT Number (if registered)", entity.vat_number),
            ("Country of Incorporation and Domicile", entity.country_of_incorporation),
            ("Nature of Business and Principal Activities", entity.nature_of_business),
            ("Registered Office Address", entity.registered_office_address),
            ("Business Address", entity.business_address),
            ("Postal Address", entity.postal_address),
            ("Bankers", entity.bankers),
            None,
            ("PRACTITIONER / COMPILER",),
            ("Practitioner Name", entity.practitioner_name),
            ("Practitioner Firm", entity.practitioner_firm),
            ("Report Type", entity.report_type.value),
            None,
            ("REPORTING BASIS",),
            ("Financial Year End (current year)", fy.year_end_date),
            ("Financial Year End (prior/comparative year)", fy.comparative_year_end_date),
            ("Date Annual Financial Statements Approved", fy.date_approved),
            ("IFRS for SME's Edition Applied", fy.ifrs_edition.value),
            ("Company Qualifies as a Small Business Corporation (SBC) for tax", "Yes" if entity.is_sbc else "No"),
            None,
            ("OPENING BALANCES (start of the COMPARATIVE/prior year - needed for the 3-row Statement of Changes in Equity)",),
            ("Retained Income - Opening Balance (start of prior year)", fy.opening_retained_income),
            ("Share Capital - Opening Balance (start of prior year)", fy.opening_share_capital),
            ("Cash and Cash Equivalents - Opening Balance (start of prior year)", fy.opening_cash),
            None,
            ("PRIOR YEAR SUPPLEMENTARY FIGURE (validation only)",),
            ("Prior Year Depreciation Charge (total, all asset classes)", fy.prior_year_depreciation_charge),
        ],
    )

    ws = wb.create_sheet("Directors")
    ws.cell(row=1, column=2, value="Directors").font = Font(bold=True, size=14)
    headers = ["Full Name", "Nationality", "Date Appointed", "Date Resigned (if any)", "Signs Approval? (Y/N)", "Notes"]
    for i, h in enumerate(headers):
        ws.cell(row=4, column=2 + i, value=h).font = BOLD
    for r, d in enumerate(entity.directors, start=5):
        ws.cell(row=r, column=2, value=d.full_name)
        ws.cell(row=r, column=3, value=d.nationality)
        ws.cell(row=r, column=4, value=d.date_appointed)
        ws.cell(row=r, column=5, value=d.date_resigned)
        ws.cell(row=r, column=6, value="Y" if d.signs_approval else "N")
        ws.cell(row=r, column=7, value=d.notes)

    ws = wb.create_sheet("Trial Balance")
    ws.cell(row=1, column=2, value="Trial Balance").font = Font(bold=True, size=14)
    headers = ["Account Name", "AFS Category", "Current Year Amount", "Prior Year Amount", "Notes"]
    for i, h in enumerate(headers):
        ws.cell(row=4, column=2 + i, value=h).font = BOLD
    for r, line in enumerate(fy.trial_balance_lines, start=5):
        ws.cell(row=r, column=2, value=line.account_name)
        ws.cell(row=r, column=3, value=line.afs_category.value)
        ws.cell(row=r, column=4, value=line.current_year_amount)
        ws.cell(row=r, column=5, value=line.prior_year_amount)
        ws.cell(row=r, column=6, value=line.notes)

    ws = wb.create_sheet("PPE Register")
    ws.cell(row=1, column=2, value="PPE Register").font = Font(bold=True, size=14)
    headers = [
        "Asset Category", "Depreciation Method", "Useful Life (yrs)", "Opening Cost", "Additions",
        "Disposals - Cost", "Closing Cost", "Opening Accum. Depreciation", "Depreciation Charge",
        "Accum. Depreciation on Disposals", "Closing Accum. Depreciation", "Carrying Value",
    ]
    for i, h in enumerate(headers):
        ws.cell(row=4, column=2 + i, value=h).font = BOLD
    for r, a in enumerate(fy.ppe_assets, start=5):
        ws.cell(row=r, column=2, value=a.asset_category)
        ws.cell(row=r, column=3, value=a.depreciation_method)
        ws.cell(row=r, column=4, value=a.useful_life_years)
        ws.cell(row=r, column=5, value=a.opening_cost)
        ws.cell(row=r, column=6, value=a.additions)
        ws.cell(row=r, column=7, value=a.disposals_cost)
        ws.cell(row=r, column=8, value=a.closing_cost)
        ws.cell(row=r, column=9, value=a.opening_accumulated_depreciation)
        ws.cell(row=r, column=10, value=a.depreciation_charge)
        ws.cell(row=r, column=11, value=a.accumulated_depreciation_on_disposals)
        ws.cell(row=r, column=12, value=a.closing_accumulated_depreciation)
        ws.cell(row=r, column=13, value=a.carrying_value)

    ws = wb.create_sheet("Loans to-from Shareholders")
    ws.cell(row=1, column=2, value="Loans to-from Shareholders").font = Font(bold=True, size=14)
    headers = [
        "Shareholder / Director Name", "Direction\n(To / From)", "Opening Balance", "Advances", "Repayments",
        "Interest Rate % p.a.", "Interest Charged\n(received in cash)", "Closing Balance",
        "Secured / Unsecured", "Repayment Terms",
    ]
    for i, h in enumerate(headers):
        ws.cell(row=4, column=2 + i, value=h).font = BOLD
    for r, l in enumerate(fy.shareholder_loans, start=5):
        ws.cell(row=r, column=2, value=l.shareholder_name)
        ws.cell(row=r, column=3, value=l.direction.value)
        ws.cell(row=r, column=4, value=l.opening_balance)
        ws.cell(row=r, column=5, value=l.advances)
        ws.cell(row=r, column=6, value=l.repayments)
        ws.cell(row=r, column=7, value=l.interest_rate_pa)
        ws.cell(row=r, column=8, value=l.interest_charged)
        ws.cell(row=r, column=9, value=l.closing_balance)
        ws.cell(row=r, column=10, value=l.secured_or_unsecured)
        ws.cell(row=r, column=11, value=l.repayment_terms)

    ws = wb.create_sheet("Tax Computation")
    ws.cell(row=1, column=2, value="Tax Computation").font = Font(bold=True, size=14)
    ws.cell(row=4, column=2, value="Net profit per income statement")
    ws.cell(row=4, column=3, value=report.tax.net_profit_per_income_statement)
    ws.cell(row=6, column=2, value="Temporary differences").font = BOLD
    ws.cell(row=7, column=2, value="Description").font = BOLD
    ws.cell(row=7, column=3, value="Amount").font = BOLD
    r = 8
    for td in fy.tax_differences:
        ws.cell(row=r, column=2, value=td.description)
        ws.cell(row=r, column=3, value=td.amount)
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Total temporary differences")
    ws.cell(row=r, column=3, value=report.tax.total_temporary_differences)
    r += 2
    ws.cell(row=r, column=2, value="Calculated tax profit for the year")
    ws.cell(row=r, column=3, value=report.tax.calculated_tax_profit)
    r += 1
    ws.cell(row=r, column=2, value="Assessed loss brought forward")
    ws.cell(row=r, column=3, value=fy.tax_computation_meta.assessed_loss_brought_forward if fy.tax_computation_meta else 0)
    r += 1
    ws.cell(row=r, column=2, value="Assessed loss utilised")
    ws.cell(row=r, column=3, value=report.tax.assessed_loss_utilised)
    r += 1
    ws.cell(row=r, column=2, value="Taxable income for the year")
    ws.cell(row=r, column=3, value=report.tax.taxable_income)
    r += 1
    ws.cell(row=r, column=2, value="Assessed loss carried forward")
    ws.cell(row=r, column=3, value=report.tax.assessed_loss_carried_forward)
    r += 2
    ws.cell(row=r, column=2, value="SBC / Company tax bracket table - EDIT FOR THE APPLICABLE YEAR OF ASSESSMENT - CHECK CURRENT SARS RATES").font = BOLD
    r += 1
    ws.cell(row=r, column=2, value="Lower Limit").font = BOLD
    ws.cell(row=r, column=3, value="Upper Limit").font = BOLD
    ws.cell(row=r, column=4, value="Rate").font = BOLD
    r += 1
    for b in fy.tax_brackets:
        ws.cell(row=r, column=2, value=b.lower_limit)
        ws.cell(row=r, column=3, value=b.upper_limit if b.upper_limit is not None else "N/A")
        ws.cell(row=r, column=4, value=b.rate)
        r += 1

    ws = wb.create_sheet("Policy Elections")
    ws.cell(row=1, column=2, value="Policy Elections").font = Font(bold=True, size=14)
    ws.cell(row=4, column=2, value="Accounting Policy Area (IFRS for SME's Section)").font = BOLD
    ws.cell(row=4, column=3, value="Applicable? (Y/N)").font = BOLD
    ws.cell(row=4, column=4, value="Notes").font = BOLD
    existing = {pe.policy_area: pe for pe in fy.policy_elections}
    for r, area in enumerate(POLICY_AREAS, start=5):
        pe = existing.get(area)
        ws.cell(row=r, column=2, value=area)
        ws.cell(row=r, column=3, value="Y" if (pe and pe.applicable) else "N")
        ws.cell(row=r, column=4, value=pe.notes if pe else None)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
