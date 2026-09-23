"""Imports the "AFS Compiler Input Template" workbook (Section 5 of the spec) into the database.

Column layout note: every sheet in the companion template has a blank column A
and starts its real content in column B onward - we scan by label text (column B)
rather than hardcoding row numbers, so minor row insertions in the workbook don't
break the import.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import openpyxl
from sqlalchemy.orm import Session

from app.models import (
    POLICY_AREAS,
    AFSCategory,
    Director,
    Entity,
    FinancialYear,
    IFRSEdition,
    LoanDirection,
    PPEAsset,
    PolicyElection,
    ReportType,
    ShareholderLoan,
    TaxBracket,
    TaxComputationMeta,
    TaxDifference,
    TrialBalanceLine,
)


class ImportError_(Exception):
    pass


def _to_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _yn(value) -> bool:
    return str(value).strip().upper().startswith("Y")


def _num(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, str):
        try:
            return float(value.replace(",", "").strip() or 0)
        except ValueError:
            return 0.0
    return float(value)


def _label_map(ws, label_col: int = 2, value_col: int = 3) -> dict[str, tuple[int, object]]:
    """Map every non-empty column-B label on a sheet to (row_index, column-C value)."""
    out: dict[str, tuple[int, object]] = {}
    for row in ws.iter_rows():
        label_cell = row[label_col - 1]
        if label_cell.value is None:
            continue
        label = str(label_cell.value).strip()
        value_cell = row[value_col - 1] if len(row) >= value_col else None
        out[label] = (label_cell.row, value_cell.value if value_cell is not None else None)
    return out


def _find_row(ws, text: str, col: int = 2) -> int | None:
    for row in ws.iter_rows():
        cell = row[col - 1]
        if cell.value is not None and str(cell.value).strip() == text:
            return cell.row
    return None


def parse_entity_setup(ws) -> dict:
    labels = _label_map(ws)

    def g(label):
        entry = labels.get(label)
        return entry[1] if entry else None

    return {
        "company_name": g("Company Name"),
        "registration_number": g("Registration Number"),
        "tax_reference_number": g("Tax Reference Number"),
        "vat_number": g("VAT Number (if registered)"),
        "country_of_incorporation": g("Country of Incorporation and Domicile") or "South Africa",
        "nature_of_business": g("Nature of Business and Principal Activities"),
        "registered_office_address": g("Registered Office Address"),
        "business_address": g("Business Address"),
        "postal_address": g("Postal Address"),
        "bankers": g("Bankers"),
        "practitioner_name": g("Practitioner Name"),
        "practitioner_firm": g("Practitioner Firm"),
        "report_type": g("Report Type") or ReportType.COMPILATION.value,
        "year_end_date": _to_date(g("Financial Year End (current year)")),
        "comparative_year_end_date": _to_date(g("Financial Year End (prior/comparative year)")),
        "date_approved": _to_date(g("Date Annual Financial Statements Approved")),
        "ifrs_edition": g("IFRS for SME's Edition Applied") or IFRSEdition.SECOND_2015.value,
        "is_sbc": _yn(g("Company Qualifies as a Small Business Corporation (SBC) for tax") or "Y"),
        "opening_retained_income": _num(g("Retained Income - Opening Balance (start of prior year)")),
        "opening_share_capital": _num(g("Share Capital - Opening Balance (start of prior year)")),
        "opening_cash": _num(g("Cash and Cash Equivalents - Opening Balance (start of prior year)")),
        "prior_year_depreciation_charge": _num(g("Prior Year Depreciation Charge (total, all asset classes)")),
    }


def parse_directors(ws) -> list[dict]:
    header_row = _find_row(ws, "Full Name")
    if header_row is None:
        return []
    out = []
    for row in ws.iter_rows(min_row=header_row + 1):
        name = row[1].value  # column B
        if not name:
            continue
        out.append(
            {
                "full_name": str(name).strip(),
                "nationality": row[2].value,
                "date_appointed": _to_date(row[3].value),
                "date_resigned": _to_date(row[4].value),
                "signs_approval": _yn(row[5].value) if row[5].value is not None else True,
                "notes": row[6].value if len(row) > 6 else None,
            }
        )
    return out


def parse_trial_balance(ws) -> list[dict]:
    header_row = _find_row(ws, "Account Name")
    if header_row is None:
        return []
    valid_categories = {c.value for c in AFSCategory}
    out = []
    for row in ws.iter_rows(min_row=header_row + 1):
        account_name = row[1].value
        if not account_name:
            continue
        if str(account_name).strip().upper().startswith("CATEGORY SUBTOTALS"):
            break
        category = row[2].value
        if category not in valid_categories:
            continue
        out.append(
            {
                "account_name": str(account_name).strip(),
                "afs_category": category,
                "current_year_amount": _num(row[3].value),
                "prior_year_amount": _num(row[4].value),
                "notes": row[5].value if len(row) > 5 else None,
            }
        )
    return out


def parse_ppe_register(ws) -> list[dict]:
    header_row = _find_row(ws, "Asset Category")
    if header_row is None:
        return []
    out = []
    for row in ws.iter_rows(min_row=header_row + 1):
        category = row[1].value
        if not category or str(category).strip().lower() == "total":
            continue
        out.append(
            {
                "asset_category": str(category).strip(),
                "depreciation_method": row[2].value or "Straight line",
                "useful_life_years": _num(row[3].value),
                "opening_cost": _num(row[4].value),
                "additions": _num(row[5].value),
                "disposals_cost": _num(row[6].value),
                # column H (index 7 -> row[7]) is the computed Closing Cost - skip.
                "opening_accumulated_depreciation": _num(row[8].value),
                "depreciation_charge": _num(row[9].value),
                "accumulated_depreciation_on_disposals": _num(row[10].value),
                # column L/M are computed closing accum. dep. / carrying value - skip.
            }
        )
    return out


def parse_shareholder_loans(ws) -> list[dict]:
    header_row = _find_row(ws, "Shareholder / Director Name")
    if header_row is None:
        return []
    out = []
    for row in ws.iter_rows(min_row=header_row + 1):
        name = row[1].value
        if not name or str(name).strip().lower() == "total":
            continue
        direction = str(row[2].value or "To").strip()
        if direction not in (LoanDirection.TO.value, LoanDirection.FROM.value):
            direction = LoanDirection.TO.value
        out.append(
            {
                "shareholder_name": str(name).strip(),
                "direction": direction,
                "opening_balance": _num(row[3].value),
                "advances": _num(row[4].value),
                "repayments": _num(row[5].value),
                "interest_rate_pa": _num(row[6].value),
                "interest_charged": _num(row[7].value),
                # column I (index 8) is the computed Closing Balance - skip.
                "secured_or_unsecured": row[9].value or "Unsecured",
                "repayment_terms": row[10].value,
            }
        )
    return out


def parse_tax_computation(ws) -> dict:
    labels = _label_map(ws)
    assessed_loss_brought_forward = _num(
        labels.get("Assessed loss brought forward", (None, 0))[1]
    )

    # Temporary differences: rows between "Description"/"Amount" header and "Total temporary differences".
    desc_header_row = None
    for row in ws.iter_rows():
        if row[1].value == "Description" and row[2].value == "Amount":
            desc_header_row = row[1].row
            break

    differences = []
    if desc_header_row is not None:
        for row in ws.iter_rows(min_row=desc_header_row + 1):
            desc = row[1].value
            if desc is None:
                break
            if str(desc).strip().lower().startswith("total temporary differences"):
                break
            amount = row[2].value
            if amount is None:
                continue
            differences.append({"description": str(desc).strip(), "amount": _num(amount)})

    # Active tax bracket table: the FIRST "Lower Limit"/"Upper Limit"/"Rate" header found
    # (the "Reference only" table further down is for a future year - ignored).
    bracket_header_row = None
    for row in ws.iter_rows():
        if row[1].value == "Lower Limit" and row[2].value == "Upper Limit":
            bracket_header_row = row[1].row
            break

    brackets = []
    if bracket_header_row is not None:
        for row in ws.iter_rows(min_row=bracket_header_row + 1):
            lower = row[1].value
            if lower is None:
                break
            upper = row[2].value
            rate = row[3].value
            brackets.append(
                {
                    "lower_limit": _num(lower),
                    "upper_limit": None if (upper is None or str(upper).strip().upper() == "N/A") else _num(upper),
                    "rate": _num(rate),
                }
            )

    return {
        "assessed_loss_brought_forward": assessed_loss_brought_forward,
        "temporary_differences": differences,
        "tax_brackets": brackets,
    }


def parse_policy_elections(ws) -> list[dict]:
    header_row = _find_row(ws, "Accounting Policy Area (IFRS for SME's Section)")
    out = []
    if header_row is None:
        return [{"policy_area": area, "applicable": False, "notes": None} for area in POLICY_AREAS]
    seen = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        area = row[1].value
        if not area:
            continue
        seen[str(area).strip()] = {
            "policy_area": str(area).strip(),
            "applicable": _yn(row[2].value) if row[2].value is not None else False,
            "notes": row[3].value if len(row) > 3 else None,
        }
    # Preserve canonical order; include any areas the workbook has that we don't recognise too.
    ordered = [seen[a] for a in POLICY_AREAS if a in seen]
    ordered += [v for k, v in seen.items() if k not in POLICY_AREAS]
    return ordered


def import_workbook(db: Session, file_bytes: bytes, entity_id: int | None = None) -> FinancialYear:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)

    entity_data = parse_entity_setup(wb["Entity Setup"])
    if not entity_data["company_name"] or not entity_data["registration_number"]:
        raise ImportError_("Entity Setup sheet is missing Company Name or Registration Number.")
    if entity_data["year_end_date"] is None:
        raise ImportError_("Entity Setup sheet is missing the current-year Financial Year End date.")

    directors_data = parse_directors(wb["Directors"])
    tb_data = parse_trial_balance(wb["Trial Balance"])
    ppe_data = parse_ppe_register(wb["PPE Register"])
    loans_data = parse_shareholder_loans(wb["Loans to-from Shareholders"])
    tax_data = parse_tax_computation(wb["Tax Computation"])
    policy_data = parse_policy_elections(wb["Policy Elections"])

    # --- Entity: find-or-create by explicit id, else by registration number ---
    entity = None
    if entity_id is not None:
        entity = db.get(Entity, entity_id)
    if entity is None:
        entity = db.query(Entity).filter_by(registration_number=entity_data["registration_number"]).one_or_none()
    if entity is None:
        entity = Entity(registration_number=entity_data["registration_number"])
        db.add(entity)

    entity.company_name = entity_data["company_name"]
    entity.registration_number = entity_data["registration_number"]
    entity.tax_reference_number = entity_data["tax_reference_number"]
    entity.vat_number = entity_data["vat_number"]
    entity.country_of_incorporation = entity_data["country_of_incorporation"]
    entity.nature_of_business = entity_data["nature_of_business"]
    entity.registered_office_address = entity_data["registered_office_address"]
    entity.business_address = entity_data["business_address"]
    entity.postal_address = entity_data["postal_address"]
    entity.bankers = entity_data["bankers"]
    entity.practitioner_name = entity_data["practitioner_name"]
    entity.practitioner_firm = entity_data["practitioner_firm"]
    entity.report_type = ReportType(entity_data["report_type"])
    entity.is_sbc = entity_data["is_sbc"]

    db.flush()

    # Directors: full replace from the workbook.
    for d in list(entity.directors):
        db.delete(d)
    db.flush()
    for d in directors_data:
        db.add(Director(entity_id=entity.id, **d))

    # --- FinancialYear: find-or-create for this entity + year-end date ---
    fy = (
        db.query(FinancialYear)
        .filter_by(entity_id=entity.id, year_end_date=entity_data["year_end_date"])
        .one_or_none()
    )
    if fy is None:
        fy = FinancialYear(entity_id=entity.id, year_end_date=entity_data["year_end_date"])
        db.add(fy)
    else:
        for collection in (
            fy.trial_balance_lines,
            fy.ppe_assets,
            fy.shareholder_loans,
            fy.tax_differences,
            fy.tax_brackets,
            fy.policy_elections,
        ):
            for item in list(collection):
                db.delete(item)
        if fy.tax_computation_meta is not None:
            db.delete(fy.tax_computation_meta)
        db.flush()

    fy.comparative_year_end_date = entity_data["comparative_year_end_date"]
    fy.date_approved = entity_data["date_approved"]
    fy.ifrs_edition = IFRSEdition(entity_data["ifrs_edition"])
    fy.opening_retained_income = entity_data["opening_retained_income"]
    fy.opening_share_capital = entity_data["opening_share_capital"]
    fy.opening_cash = entity_data["opening_cash"]
    fy.prior_year_depreciation_charge = entity_data["prior_year_depreciation_charge"]

    db.flush()

    for line in tb_data:
        db.add(TrialBalanceLine(financial_year_id=fy.id, **line))
    for asset in ppe_data:
        db.add(PPEAsset(financial_year_id=fy.id, **asset))
    for loan in loans_data:
        db.add(ShareholderLoan(financial_year_id=fy.id, **loan))

    db.add(
        TaxComputationMeta(
            financial_year_id=fy.id,
            assessed_loss_brought_forward=tax_data["assessed_loss_brought_forward"],
        )
    )
    for diff in tax_data["temporary_differences"]:
        db.add(TaxDifference(financial_year_id=fy.id, **diff))
    for bracket in tax_data["tax_brackets"]:
        db.add(TaxBracket(financial_year_id=fy.id, **bracket))

    for policy in policy_data:
        db.add(PolicyElection(financial_year_id=fy.id, **policy))

    db.commit()
    db.refresh(fy)
    return fy


def existing_financial_year(db: Session, entity_id: int, year_end_date) -> FinancialYear | None:
    return db.query(FinancialYear).filter_by(entity_id=entity_id, year_end_date=year_end_date).one_or_none()


def peek_workbook_identity(file_bytes: bytes) -> dict:
    """Cheap pre-read used to show the confirmation prompt before a destructive overwrite."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    entity_data = parse_entity_setup(wb["Entity Setup"])
    return entity_data
